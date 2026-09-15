import argparse
import re
import sys
import unicodedata
from pathlib import Path

import requests

try:
    import pdfplumber
except ImportError:
    pdfplumber = None

try:
    import docx
except ImportError:
    docx = None


# ============================================================
# CONFIGURACIÓN
# ============================================================

EXTENSIONES_VALIDAS = {
    ".pdf",
    ".docx",
}


# ============================================================
# NORMALIZACIÓN
# ============================================================

def normalizar_nombre(texto: str) -> str:
    if not texto:
        return ""

    texto = texto.strip().lower()

    texto = unicodedata.normalize(
        "NFKD",
        texto,
    )

    texto = "".join(
        caracter
        for caracter in texto
        if not unicodedata.combining(caracter)
    )

    texto = re.sub(
        r"[^a-z0-9]+",
        " ",
        texto,
    )

    texto = re.sub(
        r"\s+",
        " ",
        texto,
    )

    return texto.strip()


def limpiar_texto(texto: str) -> str:
    if not texto:
        return ""

    texto = texto.replace(
        "\r",
        "\n",
    )

    texto = re.sub(
        r"[ \t]+",
        " ",
        texto,
    )

    texto = re.sub(
        r"\n{3,}",
        "\n\n",
        texto,
    )

    return texto.strip()


# ============================================================
# PDF
# ============================================================

def extraer_texto_pdf(ruta: Path) -> str:

    if pdfplumber is None:

        print(
            f"⚠ pdfplumber no está instalado. "
            f"Se omite: {ruta.name}"
        )

        return ""

    partes = []

    try:

        with pdfplumber.open(ruta) as pdf:

            for pagina in pdf.pages:

                texto = pagina.extract_text()

                if texto:

                    partes.append(
                        texto.strip()
                    )

    except Exception as error:

        print(
            f"⚠ Error leyendo PDF "
            f"{ruta.name}: {error}"
        )

    return "\n".join(partes)


# ============================================================
# WORD
# ============================================================

def extraer_texto_docx(ruta: Path) -> str:

    if docx is None:

        print(
            f"⚠ python-docx no está instalado. "
            f"Se omite: {ruta.name}"
        )

        return ""

    try:

        documento = docx.Document(ruta)

        partes = []

        # ----------------------------------------------------
        # PÁRRAFOS
        # ----------------------------------------------------

        for parrafo in documento.paragraphs:

            texto = parrafo.text.strip()

            if texto:

                partes.append(texto)

        # ----------------------------------------------------
        # TABLAS
        # ----------------------------------------------------

        for tabla in documento.tables:

            for fila in tabla.rows:

                celdas = []

                for celda in fila.cells:

                    texto = celda.text.strip()

                    if texto:

                        celdas.append(texto)

                if celdas:

                    partes.append(
                        " | ".join(celdas)
                    )

        return "\n".join(partes)

    except Exception as error:

        print(
            f"⚠ Error leyendo Word "
            f"{ruta.name}: {error}"
        )

        return ""


# ============================================================
# ARCHIVOS
# ============================================================

def extraer_texto_archivo(ruta: Path) -> str:

    extension = ruta.suffix.lower()

    if extension == ".pdf":
        return extraer_texto_pdf(ruta)

    if extension == ".docx":
        return extraer_texto_docx(ruta)

    return ""


def buscar_documentos_actividad(
    carpeta_actividad: Path,
) -> list[Path]:

    archivos = []

    for archivo in carpeta_actividad.rglob("*"):

        if (
            archivo.is_file()
            and archivo.suffix.lower()
            in EXTENSIONES_VALIDAS
        ):

            archivos.append(archivo)

    return archivos


# ============================================================
# TEMPLATE VACÍO
# ============================================================

def es_checklist_sin_informacion(
    texto: str,
) -> bool:

    if not texto:
        return True

    texto_lower = texto.lower()

    indicadores = [
        "checklist research",
        "a quién pedir la info",
        "argentina: tati",
        "chile: nacha",
        "colombia - perú - bolivia",
        "research template",
        "uso interno",
    ]

    cantidad = sum(
        1
        for indicador in indicadores
        if indicador in texto_lower
    )

    if cantidad < 3:
        return False

    texto_prueba = texto_lower

    frases_internas = [
        "checklist research",
        "actividades",
        "actividad:",
        "responsable:",
        "recurso / contenido",
        "link / archivo",
        "recursos",
        "contenido",
        "buscar",
        "página web del to",
        "link civitatis",
        "link formulario de carga",
        "brochure / pdf",
        "id servicio de reserva",
        "descripción principal",
        "highlights",
        "itinerario",
        "mínimo 4 imágenes de calidad",
        "horarios",
        "restricciones",
        "obligatorio",
        "notas adicionales",
        "a quién pedir la info",
        "argentina: tati",
        "colombia - perú - bolivia: mafe",
        "brasil - méxico - costa rica: alice",
        "chile: nacha",
        "tur.com",
        "research template",
        "uso interno",
    ]

    for frase in frases_internas:

        texto_prueba = texto_prueba.replace(
            frase,
            "",
        )

    texto_prueba = re.sub(
        r"[\W_]+",
        "",
        texto_prueba,
    )

    return len(texto_prueba) < 100


# ============================================================
# EXTRAER SECCIONES
# ============================================================

def linea_es_encabezado(
    linea: str,
    titulo: str,
) -> bool:

    linea_lower = linea.strip().lower()
    titulo_lower = titulo.lower()

    return (
        linea_lower == titulo_lower
        or linea_lower.startswith(
            titulo_lower + " —"
        )
        or linea_lower.startswith(
            titulo_lower + " -"
        )
        or linea_lower.startswith(
            titulo_lower + " |"
        )
        or linea_lower.startswith(
            titulo_lower + ":"
        )
    )


def extraer_seccion(
    texto: str,
    inicio: str,
    siguientes: list[str],
) -> str | None:

    if not texto:
        return None

    lineas = texto.splitlines()

    indice_inicio = None
    contenido_misma_linea = None

    for i, linea in enumerate(lineas):

        linea_limpia = linea.strip()
        linea_lower = linea_limpia.lower()

        # ----------------------------------------------------
        # Evitar confundir "Descripción principal"
        # del checklist con la descripción real
        # ----------------------------------------------------

        if inicio.lower() == "descripción":

            es_valida = (
                linea_lower == "descripción"
                or linea_lower.startswith(
                    "descripción —"
                )
                or linea_lower.startswith(
                    "descripción -"
                )
                or linea_lower.startswith(
                    "descripción |"
                )
                or linea_lower.startswith(
                    "descripción:"
                )
            )

            if not es_valida:
                continue

        else:

            if not linea_es_encabezado(
                linea_limpia,
                inicio,
            ):
                continue

        indice_inicio = i

        # ----------------------------------------------------
        # CONTENIDO EN LA MISMA CELDA/FILA
        # ----------------------------------------------------

        if "|" in linea_limpia:

            partes = linea_limpia.split("|")

            if len(partes) >= 2:

                posible = "|".join(
                    partes[1:]
                ).strip()

                if posible:
                    contenido_misma_linea = posible

        break

    if indice_inicio is None:
        return None

    contenido = []

    if contenido_misma_linea:
        contenido.append(
            contenido_misma_linea
        )

    for linea in lineas[
        indice_inicio + 1:
    ]:

        linea_limpia = linea.strip()

        detener = False

        for siguiente in siguientes:

            if linea_es_encabezado(
                linea_limpia,
                siguiente,
            ):

                detener = True
                break

        if detener:
            break

        if linea_limpia:

            contenido.append(
                linea_limpia
            )

    resultado = "\n".join(
        contenido
    ).strip()

    if not resultado:
        return None

    return limpiar_texto(resultado)


# ============================================================
# DESCRIPCIÓN
# ============================================================

def limpiar_linea_actividad(
    linea: str,
    nombre_actividad: str,
) -> str | None:

    linea_limpia = linea.strip()

    linea_limpia = re.sub(
        r"^\s*actividad\s*:\s*",
        "",
        linea_limpia,
        flags=re.IGNORECASE,
    )

    # Si después de eliminar "Actividad:"
    # solo queda repetido el nombre de la actividad,
    # se elimina completamente.
    if (
        normalizar_nombre(linea_limpia)
        == normalizar_nombre(
            nombre_actividad
        )
    ):

        return None

    return linea_limpia


def limpiar_descripcion_final(
    descripcion: str | None,
    nombre_actividad: str,
) -> str | None:

    if not descripcion:
        return None

    texto = limpiar_texto(descripcion)

    if not texto:
        return None

    ignorar = {
        "buscar",
        "obligatorio",
        "principal",
        "principal | obligatorio",
        "principal | obligatorio | obligatorio",
        "qué es, dónde, duración, para quién",
        "que es, donde, duracion, para quien",
        "descripción de la experiencia",
        "descripcion de la experiencia",
    }

    lineas_validas = []

    for linea in texto.splitlines():

        linea_limpia = linea.strip()

        if not linea_limpia:
            continue

        if linea_limpia.lower() in ignorar:
            continue

        linea_limpia = limpiar_linea_actividad(
            linea_limpia,
            nombre_actividad,
        )

        if linea_limpia:

            lineas_validas.append(
                linea_limpia
            )

    resultado = "\n".join(
        lineas_validas
    ).strip()

    if not resultado:
        return None

    if len(resultado) < 10:
        return None

    return resultado


def extraer_descripcion_completa(
    descripcion: str | None,
    nombre_actividad: str,
) -> str | None:

    if not descripcion:
        return None

    texto = limpiar_texto(
        descripcion
    )

    lineas = texto.splitlines()

    indice_corta = None
    indice_larga = None

    for i, linea in enumerate(lineas):

        normalizada = (
            linea
            .strip()
            .lower()
            .rstrip(":")
        )

        if normalizada == "corta":
            indice_corta = i

        elif normalizada == "larga":
            indice_larga = i

    descripcion_corta = None
    descripcion_larga = None

    # --------------------------------------------------------
    # CORTA + LARGA
    # --------------------------------------------------------

    if (
        indice_corta is not None
        and indice_larga is not None
        and indice_corta < indice_larga
    ):

        descripcion_corta = "\n".join(
            lineas[
                indice_corta + 1:
                indice_larga
            ]
        ).strip()

        descripcion_larga = "\n".join(
            lineas[
                indice_larga + 1:
            ]
        ).strip()

    # --------------------------------------------------------
    # SOLO CORTA
    # --------------------------------------------------------

    elif indice_corta is not None:

        descripcion_corta = "\n".join(
            lineas[
                indice_corta + 1:
            ]
        ).strip()

    # --------------------------------------------------------
    # SOLO LARGA
    # --------------------------------------------------------

    elif indice_larga is not None:

        descripcion_larga = "\n".join(
            lineas[
                indice_larga + 1:
            ]
        ).strip()

    # --------------------------------------------------------
    # DESCRIPCIÓN ÚNICA
    # --------------------------------------------------------

    else:

        descripcion_corta = texto

    descripcion_corta = (
        limpiar_descripcion_final(
            descripcion_corta,
            nombre_actividad,
        )
    )

    descripcion_larga = (
        limpiar_descripcion_final(
            descripcion_larga,
            nombre_actividad,
        )
    )

    partes = [
        parte
        for parte in [
            descripcion_corta,
            descripcion_larga,
        ]
        if parte
    ]

    if not partes:
        return None

    return "\n\n".join(partes)


def extraer_descripcion_fallback(
    texto: str,
    nombre_actividad: str,
) -> str | None:

    """
    Se usa si no existe la sección estándar
    'Descripción — ...'.

    Busca encabezados alternativos comunes.
    """

    encabezados = [
        "Descripción de la experiencia",
        "Descripcion de la experiencia",
    ]

    siguientes = [
        "Highlights",
        "Lo más destacado",
        "Itinerario",
        "Imágenes",
        "Horarios",
        "Restricciones",
        "Incluye",
        "No incluye",
        "Notas adicionales",
    ]

    for encabezado in encabezados:

        resultado = extraer_seccion(
            texto,
            encabezado,
            siguientes,
        )

        resultado = limpiar_descripcion_final(
            resultado,
            nombre_actividad,
        )

        if resultado:
            return resultado

    return None


# ============================================================
# DURACIÓN
# ============================================================

def extraer_duracion(
    texto: str,
    nombre_actividad: str,
) -> str | None:

    if not texto:
        texto = ""

    combinado = (
        f"{nombre_actividad}\n{texto}"
    )

    # --------------------------------------------------------
    # EJEMPLOS:
    # 2D/1N
    # 3D-2N
    # --------------------------------------------------------

    patron_dn = re.search(
        r"\b(\d+)\s*d\s*[/\-]\s*(\d+)\s*n\b",
        combinado,
        re.IGNORECASE,
    )

    if patron_dn:

        dias = int(
            patron_dn.group(1)
        )

        noches = int(
            patron_dn.group(2)
        )

        texto_dias = (
            "día"
            if dias == 1
            else "días"
        )

        texto_noches = (
            "noche"
            if noches == 1
            else "noches"
        )

        return (
            f"{dias} {texto_dias} / "
            f"{noches} {texto_noches}"
        )

    # --------------------------------------------------------
    # 2 días / 1 noche
    # 2 días y 1 noche
    # --------------------------------------------------------

    patron_dias_noches = re.search(
        r"\b(\d+)\s*d[ií]as?"
        r"\s*(?:/|y|-)\s*"
        r"(\d+|una?)\s*noches?\b",
        combinado,
        re.IGNORECASE,
    )

    if patron_dias_noches:

        dias = int(
            patron_dias_noches.group(1)
        )

        noches_texto = (
            patron_dias_noches.group(2)
            .lower()
        )

        noches = (
            1
            if noches_texto
            in {"un", "una"}
            else int(noches_texto)
        )

        return (
            f"{dias} "
            f"{'día' if dias == 1 else 'días'}"
            f" / "
            f"{noches} "
            f"{'noche' if noches == 1 else 'noches'}"
        )

    # --------------------------------------------------------
    # DURACIÓN EXPLÍCITA
    # --------------------------------------------------------

    patrones = [

        r"duraci[oó]n\s+de\s+la\s+actividad"
        r"\s*[:\-]?\s*"
        r"(\d+(?:[.,]\d+)?\s*minutos?)",

        r"duraci[oó]n\s+de\s+la\s+actividad"
        r"\s*[:\-]?\s*"
        r"(\d+(?:[.,]\d+)?\s*horas?)",

        r"duraci[oó]n(?:\s+total)?"
        r"\s*(?:de)?\s*[:\-]?\s*"
        r"(\d+(?:[.,]\d+)?\s*horas?)",

        r"duraci[oó]n(?:\s+total)?"
        r"\s*(?:de)?\s*[:\-]?\s*"
        r"(\d+(?:[.,]\d+)?\s*hrs?)",

        r"duraci[oó]n(?:\s+total)?"
        r"\s*(?:de)?\s*[:\-]?\s*"
        r"(\d+(?:[.,]\d+)?\s*h\b)",

        r"duraci[oó]n(?:\s+total)?"
        r"\s*(?:de)?\s*[:\-]?\s*"
        r"(\d+\s*minutos?)",

        r"duraci[oó]n(?:\s+total)?"
        r"\s*(?:de)?\s*[:\-]?\s*"
        r"(\d+\s*d[ií]as?)",
    ]

    for patron in patrones:

        resultado = re.search(
            patron,
            combinado,
            re.IGNORECASE,
        )

        if resultado:

            return (
                resultado
                .group(1)
                .strip()
            )

    # --------------------------------------------------------
    # "aproximadamente 17 horas de recorrido"
    # --------------------------------------------------------

    resultado = re.search(
        r"(?:aproximadamente\s*)?"
        r"(\d+(?:[.,]\d+)?\s*horas?)"
        r"\s+de\s+recorrido",
        combinado,
        re.IGNORECASE,
    )

    if resultado:

        return (
            resultado
            .group(1)
            .strip()
        )

    # --------------------------------------------------------
    # FULL DAY
    # --------------------------------------------------------

    if re.search(
        r"\bfull\s*day\b",
        combinado,
        re.IGNORECASE,
    ):

        return "Full day"

    # --------------------------------------------------------
    # MEDIO DÍA
    # --------------------------------------------------------

    if re.search(
        r"\bmedio\s+d[ií]a\b",
        combinado,
        re.IGNORECASE,
    ):

        return "Medio día"

    return None


# ============================================================
# EDADES
# ============================================================

def extraer_edades(
    texto: str,
) -> tuple[int | None, int | None]:

    if not texto:
        return None, None

    edad_minima = None
    edad_maxima = None

    patrones_minimos = [

        r"edad\s+m[ií]nima"
        r"\s*[:\-]?\s*(\d+)",

        r"m[ií]nimo\s+(\d+)"
        r"\s*a[nñ]os",

        r"desde\s+los?\s+(\d+)"
        r"\s*a[nñ]os",

        # "No apto para niños menores de 5 años"
        r"(?:niños|niñas|menores)"
        r"\s+menores?\s+de\s+(\d+)"
        r"\s*a[nñ]os",
    ]

    patrones_maximos = [

        r"edad\s+m[aá]xima"
        r"\s*[:\-]?\s*(\d+)",

        r"m[aá]ximo\s+(\d+)"
        r"\s*a[nñ]os",

        r"hasta\s+los?\s+(\d+)"
        r"\s*a[nñ]os",
    ]

    for patron in patrones_minimos:

        resultado = re.search(
            patron,
            texto,
            re.IGNORECASE,
        )

        if resultado:

            edad_minima = int(
                resultado.group(1)
            )

            break

    for patron in patrones_maximos:

        resultado = re.search(
            patron,
            texto,
            re.IGNORECASE,
        )

        if resultado:

            edad_maxima = int(
                resultado.group(1)
            )

            break

    return (
        edad_minima,
        edad_maxima,
    )


# ============================================================
# IDIOMAS
# ============================================================

def extraer_idiomas(
    texto: str,
) -> list[str]:

    if not texto:
        return []

    idiomas = {

        "español":
            "Español",

        "espanol":
            "Español",

        "inglés":
            "Inglés",

        "ingles":
            "Inglés",

        "portugués":
            "Portugués",

        "portugues":
            "Portugués",

        "francés":
            "Francés",

        "frances":
            "Francés",

        "alemán":
            "Alemán",

        "aleman":
            "Alemán",

        "italiano":
            "Italiano",
    }

    encontrados = []

    texto_lower = texto.lower()

    for variante, normalizado in (
        idiomas.items()
    ):

        if variante in texto_lower:

            if normalizado not in encontrados:

                encontrados.append(
                    normalizado
                )

    return encontrados


# ============================================================
# UBICACIÓN
# ============================================================

def extraer_ubicacion(
    texto: str,
) -> str | None:

    if not texto:
        return None

    patrones = [

        r"ubicaci[oó]n"
        r"\s*[:\-]\s*"
        r"([^\n|]+)",

        r"lugar"
        r"\s*[:\-]\s*"
        r"([^\n|]+)",

        r"punto\s+de\s+encuentro"
        r"\s*[:\-]\s*"
        r"([^\n|]+)",
    ]

    for patron in patrones:

        resultado = re.search(
            patron,
            texto,
            re.IGNORECASE,
        )

        if resultado:

            ubicacion = (
                resultado
                .group(1)
                .strip()
            )

            if len(ubicacion) > 2:
                return ubicacion

    return None



# ============================================================
# CAMPOS ESTRUCTURADOS
# ============================================================

def texto_a_lista(texto: str | None) -> list[str]:
    if not texto:
        return []

    basura_exacta = {
        "buscar",
        "obligatorio",
        "obligatorio | obligatorio",
        "principal",
        "principal | obligatorio",
        "principal | obligatorio | obligatorio",
        "mínimo 4 imágenes de calidad | obligatorio | obligatorio",
        "minimo 4 imagenes de calidad | obligatorio | obligatorio",
    }

    basura_fragmentos = [
        "checklist research",
        "recurso / contenido",
        "link / archivo",
        "a quién pedir la info",
        "a quien pedir la info",
    ]

    resultado = []

    for linea in limpiar_texto(texto).splitlines():
        item = re.sub(
            r"^\s*(?:[-•●▪◦*]+|\d+[.)-])\s*",
            "",
            linea,
        ).strip()

        if not item:
            continue

        normalizada = item.lower().strip()

        if normalizada in basura_exacta:
            continue

        if any(fragmento in normalizada for fragmento in basura_fragmentos):
            continue

        # Encabezados compuestos del template, por ejemplo:
        # "Highlights — qué llevar · no permitido · qué incluye..."
        if (
            normalizada.startswith("highlights")
            and (
                "qué llevar" in normalizada
                or "que llevar" in normalizada
                or "no permitido" in normalizada
                or "incluye" in normalizada
            )
        ):
            continue

        # Evita guardar encabezados/columnas del template como contenido real.
        if (
            "obligatorio" in normalizada
            and len(normalizada.split()) <= 8
        ):
            continue

        if re.fullmatch(
            r"(?:highlights?|itinerario|horarios?|restricciones?|"
            r"imágenes?|imagenes?|notas adicionales|qué llevar|que llevar|"
            r"no permitido|incluye|no incluye)\s*[:\-–—]?",
            normalizada,
            flags=re.IGNORECASE,
        ):
            continue

        resultado.append(item)

    return list(dict.fromkeys(resultado))


def extraer_primera_seccion(texto: str, titulos: list[str], siguientes: list[str]) -> str | None:
    for titulo in titulos:
        valor = extraer_seccion(texto, titulo, siguientes)
        if valor:
            return valor
    return None


def extraer_destino(texto: str, ubicacion: str | None) -> str | None:
    if not texto:
        return ubicacion

    patrones = [
        r"(?:destino|ciudad|localidad)\s*[:\-|]\s*([^\n|]+)",
        r"(?:lugar\s+de\s+la\s+actividad)\s*[:\-|]\s*([^\n|]+)",
    ]
    for patron in patrones:
        m = re.search(patron, texto, re.IGNORECASE)
        if m:
            valor = m.group(1).strip(" .;-")
            if len(valor) > 2:
                return valor

    return ubicacion


def _normalizar_encabezado(linea: str) -> str:
    valor = normalizar_nombre(linea)
    return valor.strip()


def _titulo_base_encabezado(linea: str) -> str:
    """
    Obtiene solo la parte de título de encabezados como:
      Itinerario — paradas, tiempos, inicio y fin
      Horarios — salidas, días disponibles, duración total
      Restricciones — edad, peso, salud, clima — opcional
      Imágenes — links o indicar si están en la carpeta

    También soporta filas de tabla separadas por |.
    """
    base = linea.strip()

    if "|" in base:
        base = base.split("|", 1)[0].strip()

    # Separadores usados por el template.
    base = re.split(
        r"\s*(?:—|–)\s*",
        base,
        maxsplit=1,
    )[0].strip()

    # ":" solo se considera separador cuando lo anterior parece un
    # encabezado corto; así no rompemos frases narrativas normales.
    if ":" in base:
        izquierda = base.split(":", 1)[0].strip()
        if len(izquierda.split()) <= 6:
            base = izquierda

    return _normalizar_encabezado(base)


def _tipo_encabezado(linea: str) -> str | None:
    """
    Reconoce encabezados reales del template/documento, incluyendo
    títulos con aclaraciones después de raya larga.
    """
    original = linea.strip()
    if not original:
        return None

    n_completo = _normalizar_encabezado(original)
    n_base = _titulo_base_encabezado(original)

    if not n_base:
        return None

    reglas = {
        "descripcion": {
            "descripcion",
            "descripcion principal",
            "descripcion de la experiencia",
        },
        "highlights": {
            "highlights",
            "highlight",
            "lo mas destacado",
            "destacados",
        },
        "itinerario": {
            "itinerario",
            "programa",
            "recorrido",
        },
        "horarios": {
            "horario",
            "horarios",
            "salida",
            "salidas",
            "hora de salida",
            "horario de salida",
        },
        "restricciones": {
            "restriccion",
            "restricciones",
            "requisito",
            "requisitos",
            "no permitido",
        },
        "informacion": {
            "informacion importante",
            "info importante",
        },
        "incluye": {
            "incluye",
            "que incluye",
        },
        "no_incluye": {
            "no incluye",
            "que no incluye",
        },
        "que_llevar": {
            "que llevar",
            "que llevar al tour",
        },
        "recomendaciones": {
            "recomendaciones",
            "notas adicionales",
        },
        "imagenes": {
            "imagen",
            "imagenes",
            "fotos",
        },
    }

    for tipo, nombres in reglas.items():
        if n_base in nombres:
            return tipo

    # Caso especial de encabezado compuesto del checklist:
    # "Highlights — qué llevar · no permitido · qué incluye..."
    if n_base == "highlights":
        return "highlights"

    # No clasificar frases largas aunque contengan palabras de sección.
    if len(n_completo.split()) > 12:
        return None

    return None


def _es_encabezado_template(linea: str) -> bool:
    return _tipo_encabezado(linea) is not None


def _extraer_bloque_por_tipo(texto: str, tipos_inicio: set[str]) -> list[str]:
    lineas = texto.splitlines()
    inicio = None
    misma_linea = None

    for i, linea in enumerate(lineas):
        limpia = linea.strip()
        tipo = _tipo_encabezado(limpia)

        if tipo not in tipos_inicio:
            continue

        inicio = i

        if "|" in limpia:
            partes = [p.strip() for p in limpia.split("|")]
            if len(partes) >= 2:
                candidato = " | ".join(partes[1:]).strip()
                if candidato:
                    misma_linea = candidato
        break

    if inicio is None:
        return []

    contenido = []
    if misma_linea:
        contenido.append(misma_linea)

    for linea in lineas[inicio + 1:]:
        limpia = linea.strip()
        if not limpia:
            continue

        if _tipo_encabezado(limpia) is not None:
            break

        contenido.append(limpia)

    return texto_a_lista("\n".join(contenido))


def extraer_horarios_lista(texto: str) -> list[str]:
    """
    Prioriza una sección Horarios/Salidas real. Si no existe, toma
    únicamente líneas con forma claramente horaria; nunca bloques
    completos del itinerario.
    """
    seccion = _extraer_bloque_por_tipo(texto, {"horarios"})

    if seccion:
        filtrados = []
        for item in seccion:
            n = normalizar_nombre(item)

            if n in {
                "horario", "horarios", "salida", "salidas",
                "hora de salida", "horario de salida",
            }:
                continue

            # Dentro de una sección explícita aceptamos días/frecuencia,
            # duración y horas, pero evitamos párrafos narrativos largos.
            es_temporal = bool(re.search(
                r"(?:\b(?:[01]?\d|2[0-3])[:.]\d{2}\b|"
                r"\bcada\s+\d+\s*(?:min|minutos?|horas?)\b|"
                r"\btodos\s+los\s+d[ií]as\b|"
                r"\bd[ií]as?\s+en\s+que\b|"
                r"\bduraci[oó]n\s+de\s+la\s+actividad\b|"
                r"\bhorario\s+de\s+salida\b)",
                item,
                re.IGNORECASE,
            ))

            if es_temporal:
                filtrados.append(item)

        if filtrados:
            return list(dict.fromkeys(filtrados))

    encontrados = []

    for linea in texto.splitlines():
        l = linea.strip()
        if not l or "obligatorio" in l.lower():
            continue

        # Fallback deliberadamente estricto: exige que la propia línea
        # sea un dato horario, no que simplemente mencione una hora.
        if re.match(
            r"^(?:horario(?:s)?(?:\s+de\s+salida)?|salidas?|"
            r"d[ií]as?\s+en\s+que\s+se\s+realiza|duraci[oó]n\s+de\s+la\s+actividad)"
            r"\s*[:\-]",
            l,
            re.IGNORECASE,
        ) or re.fullmatch(
            r"todos\s+los\s+d[ií]as",
            l,
            re.IGNORECASE,
        ):
            encontrados.append(l)

    return list(dict.fromkeys(encontrados))


def extraer_restricciones_lista(texto: str) -> list[str]:
    """
    Extrae solo bloques de restricciones/no permitido/requisitos e
    información importante. Se detiene antes de Incluye/No incluye,
    Qué llevar, Itinerario, Horarios, etc.
    """
    lineas = texto.splitlines()
    encontrados = []

    tipos_inicio = {"restricciones", "informacion"}

    for i, linea in enumerate(lineas):
        if _tipo_encabezado(linea) not in tipos_inicio:
            continue

        for siguiente in lineas[i + 1:]:
            item = siguiente.strip()
            if not item:
                continue

            tipo = _tipo_encabezado(item)
            if tipo is not None:
                break

            encontrados.append(item)

    # Fallback: solo frases explícitamente restrictivas, no secciones
    # completas ni "incluye/no incluye".
    patron = re.compile(
        r"\b(?:no\s+apto|no\s+permitido|prohibid[oa]s?|restricci[oó]n|"
        r"edad\s+m[ií]nima|menores?\s+de|embarazad[oa]s?|movilidad\s+reducida|"
        r"no\s+pueden|no\s+puede|debe[n]?\s+tener|"
        r"no\s+est[aá]n\s+garantizados?|no\s+est[aá]\s+garantizad[oa])\b",
        re.IGNORECASE,
    )

    for linea in lineas:
        limpia = linea.strip()
        if not limpia or "obligatorio" in limpia.lower():
            continue

        if _tipo_encabezado(limpia) in {"incluye", "no_incluye", "que_llevar"}:
            continue

        if patron.search(limpia):
            encontrados.append(limpia)

    return texto_a_lista("\n".join(encontrados))


def extraer_highlights_lista(texto: str) -> list[str]:
    lineas = texto.splitlines()
    inicio = None

    for i, linea in enumerate(lineas):
        if _tipo_encabezado(linea) == "highlights":
            inicio = i
            break

    if inicio is None:
        return []

    contenido = []

    for linea in lineas[inicio + 1:]:
        limpia = linea.strip()
        if not limpia:
            continue

        if _tipo_encabezado(limpia) is not None:
            break

        contenido.append(limpia)

    return texto_a_lista("\n".join(contenido))


def extraer_itinerario_lista(texto: str) -> list[str]:
    lineas = texto.splitlines()
    inicio = None

    for i, linea in enumerate(lineas):
        if _tipo_encabezado(linea) == "itinerario":
            inicio = i
            break

    if inicio is None:
        return []

    contenido = []

    for linea in lineas[inicio + 1:]:
        limpia = linea.strip()
        if not limpia:
            continue

        # Algunos Word repiten "Itinerario" antes del contenido.
        if _tipo_encabezado(limpia) == "itinerario":
            continue

        if _tipo_encabezado(limpia) is not None:
            break

        contenido.append(limpia)

    return texto_a_lista("\n".join(contenido))


# ============================================================
# ANALIZAR ACTIVIDAD
# ============================================================

def analizar_contenido_actividad(
    texto: str,
    nombre_actividad: str,
) -> dict:

    texto = limpiar_texto(texto)

    encabezados_siguientes = [
        "Highlights", "Lo más destacado", "Lo mas destacado", "Destacados",
        "Itinerario", "Programa", "Recorrido",
        "Imágenes", "Imagenes", "Horarios", "Horario", "Salidas",
        "Restricciones", "Requisitos", "Incluye", "No incluye",
        "Notas adicionales", "Recomendaciones",
        "Información importante", "Informacion importante",
    ]

    descripcion_original = extraer_seccion(
        texto,
        "Descripción",
        encabezados_siguientes,
    )

    descripcion = extraer_descripcion_completa(
        descripcion_original,
        nombre_actividad,
    )

    if not descripcion:
        descripcion = extraer_descripcion_fallback(
            texto,
            nombre_actividad,
        )

    ubicacion = extraer_ubicacion(texto)
    destino = extraer_destino(texto, ubicacion)

    highlights = extraer_highlights_lista(texto)
    itinerario = extraer_itinerario_lista(texto)
    horarios = extraer_horarios_lista(texto)
    restricciones = extraer_restricciones_lista(texto)

    texto_para_duracion = "\n".join(
        parte for parte in [
            descripcion,
            "\n".join(horarios),
            texto,
        ] if parte
    )

    texto_para_edades = "\n".join(
        parte for parte in [
            descripcion,
            "\n".join(restricciones),
            texto,
        ] if parte
    )

    edad_minima, edad_maxima = extraer_edades(texto_para_edades)

    return {
        "descripcion": descripcion,
        "ubicacion": ubicacion,
        "destino": destino,
        "duracion": extraer_duracion(texto_para_duracion, nombre_actividad),
        "edadMinima": edad_minima,
        "edadMaxima": edad_maxima,
        "idiomas": extraer_idiomas(texto),
        "imagenes": [],
        "highlights": highlights,
        "itinerario": itinerario,
        "horarios": horarios,
        "restricciones": restricciones,
    }


# ============================================================
# PAÍSES
# ============================================================

def cargar_paises(
    api_base: str,
) -> list[dict]:

    try:

        respuesta = requests.get(
            f"{api_base}/api/paises",
            timeout=20,
        )

        if respuesta.status_code != 200:

            print(
                "❌ No se pudieron "
                "cargar los países."
            )

            print(
                f"   HTTP "
                f"{respuesta.status_code}"
            )

            return []

        return respuesta.json()

    except requests.RequestException as error:

        print(
            "❌ Error consultando países:"
        )

        print(
            f"   {error}"
        )

        return []


def buscar_pais(
    nombre_carpeta: str,
    paises: list[dict],
) -> dict | None:

    buscado = normalizar_nombre(
        nombre_carpeta
    )

    for pais in paises:

        nombre_bd = (
            pais.get("nombre")
            or ""
        )

        if (
            normalizar_nombre(nombre_bd)
            == buscado
        ):

            return pais

    return None


# ============================================================
# OPERADORES
# ============================================================

def cargar_operadores(
    api_base: str,
) -> list[dict]:

    try:

        respuesta = requests.get(
            f"{api_base}/api/operadores",
            timeout=20,
        )

        if respuesta.status_code != 200:

            print(
                "❌ No se pudieron "
                "cargar los operadores."
            )

            return []

        return respuesta.json()

    except requests.RequestException as error:

        print(
            "❌ Error consultando "
            "operadores:"
        )

        print(
            f"   {error}"
        )

        return []


def buscar_operador(
    nombre_operador: str,
    pais_id: int,
    operadores: list[dict],
) -> dict | None:

    buscado = normalizar_nombre(
        nombre_operador
    )

    for operador in operadores:

        pais = (
            operador.get("pais")
            or {}
        )

        if pais.get("id") != pais_id:
            continue

        nombre_bd = (
            operador.get("nombre")
            or ""
        )

        if (
            normalizar_nombre(nombre_bd)
            == buscado
        ):

            return operador

    return None


def crear_operador(
    api_base: str,
    pais_id: int,
    nombre_operador: str,
) -> dict | None:

    payload = {
        "nombre":
            nombre_operador,

        "estado":
            "ACTIVO",
    }

    try:

        respuesta = requests.post(
            f"{api_base}/api/operadores",
            params={
                "paisId":
                    pais_id
            },
            json=payload,
            timeout=30,
        )

        if respuesta.status_code in (
            200,
            201,
        ):

            operador = respuesta.json()

            print(
                "✅ Operador creado:"
            )

            print(
                f"   {nombre_operador}"
            )

            print(
                f"   ID: "
                f"{operador.get('id')}"
            )

            return operador

        print(
            "❌ No se pudo crear "
            "el operador:"
        )

        print(
            f"   {nombre_operador}"
        )

        print(
            f"   HTTP "
            f"{respuesta.status_code}"
        )

        print(
            respuesta.text[:500]
        )

        return None

    except requests.RequestException as error:

        print(
            "❌ Error creando operador:"
        )

        print(
            f"   {error}"
        )

        return None


# ============================================================
# ACTIVIDAD EXISTENTE
# ============================================================

def obtener_actividad_existente(
    backend_actividades: str,
    operador_id: int,
    nombre: str,
) -> dict | None:

    try:

        respuesta = requests.get(
            (
                f"{backend_actividades}"
                "/existe"
            ),
            params={
                "operadorTuristicoId":
                    operador_id,

                "nombre":
                    nombre,
            },
            timeout=15,
        )

        if respuesta.status_code == 200:

            return respuesta.json()

        if respuesta.status_code == 404:

            return None

        print(
            "⚠ No se pudo comprobar "
            "actividad existente."
        )

        print(
            f"   HTTP "
            f"{respuesta.status_code}"
        )

        return None

    except requests.RequestException as error:

        print(
            "⚠ Error consultando "
            "actividad:"
        )

        print(
            f"   {error}"
        )

        return None


# ============================================================
# PROTEGER DATOS EXISTENTES
# ============================================================

def valor_texto_nuevo_o_existente(
    nuevo,
    existente,
):

    if nuevo is None:
        return existente

    if isinstance(nuevo, str):

        if not nuevo.strip():
            return existente

    return nuevo


def lista_nueva_o_existente(
    nueva,
    existente,
):

    if nueva:
        return nueva

    if existente:
        return existente

    return []


def construir_payload_actualizacion(
    payload_nuevo: dict,
    existente: dict,
) -> dict:

    """
    Evita que un dato None o [] de una
    extracción incompleta borre información
    que ya estaba correctamente guardada.
    """

    return {

        "nombre":
            payload_nuevo["nombre"],

        "descripcion":
            valor_texto_nuevo_o_existente(
                payload_nuevo.get(
                    "descripcion"
                ),
                existente.get(
                    "descripcion"
                ),
            ),

        "ubicacion":
            valor_texto_nuevo_o_existente(
                payload_nuevo.get(
                    "ubicacion"
                ),
                existente.get(
                    "ubicacion"
                ),
            ),

        "duracion":
            valor_texto_nuevo_o_existente(
                payload_nuevo.get(
                    "duracion"
                ),
                existente.get(
                    "duracion"
                ),
            ),

        "edadMinima":
            valor_texto_nuevo_o_existente(
                payload_nuevo.get(
                    "edadMinima"
                ),
                existente.get(
                    "edadMinima"
                ),
            ),

        "edadMaxima":
            valor_texto_nuevo_o_existente(
                payload_nuevo.get(
                    "edadMaxima"
                ),
                existente.get(
                    "edadMaxima"
                ),
            ),

        "idiomas":
            lista_nueva_o_existente(
                payload_nuevo.get(
                    "idiomas"
                ),
                existente.get(
                    "idiomas"
                ),
            ),

        "imagenes":
            lista_nueva_o_existente(
                payload_nuevo.get(
                    "imagenes"
                ),
                existente.get(
                    "imagenes"
                ),
            ),

        "destino":
            valor_texto_nuevo_o_existente(
                payload_nuevo.get(
                    "destino"
                ),
                existente.get(
                    "destino"
                ),
            ),

        "highlights":
            lista_nueva_o_existente(
                payload_nuevo.get(
                    "highlights"
                ),
                existente.get(
                    "highlights"
                ),
            ),

        "itinerario":
            lista_nueva_o_existente(
                payload_nuevo.get(
                    "itinerario"
                ),
                existente.get(
                    "itinerario"
                ),
            ),

        "horarios":
            lista_nueva_o_existente(
                payload_nuevo.get(
                    "horarios"
                ),
                existente.get(
                    "horarios"
                ),
            ),

        "restricciones":
            lista_nueva_o_existente(
                payload_nuevo.get(
                    "restricciones"
                ),
                existente.get(
                    "restricciones"
                ),
            ),

        "urlOrigen":
            payload_nuevo[
                "urlOrigen"
            ],

        # Para PUT enviamos la relación completa ya persistida cuando
        # está disponible. Esto evita diferencias con registros antiguos
        # que no aceptan correctamente una relación reducida a {"id": ...}.
        "operadorTuristico":
            (
                existente.get("operadorTuristico")
                or payload_nuevo["operadorTuristico"]
            ),

        "activo":
            True,
    }


# ============================================================
# PROCESAMIENTO PRINCIPAL
# ============================================================

def recorrer_y_procesar(
    raiz: Path,
    api_base: str,
    dry_run: bool,
    solo_actividad: str | None = None,
    solo_operador: str | None = None,
):

    backend_actividades = (
        f"{api_base}/api/actividades"
    )

    total_creadas = 0
    total_actualizadas = 0
    total_simuladas = 0
    total_operadores_creados = 0
    total_omitidas = 0
    total_error = 0

    # --------------------------------------------------------
    # CARGAR PAÍSES Y OPERADORES
    # --------------------------------------------------------

    paises = cargar_paises(
        api_base
    )

    if not paises:

        print(
            "❌ No hay países disponibles."
        )

        return

    operadores_backend = (
        cargar_operadores(
            api_base
        )
    )

    # --------------------------------------------------------
    # PAÍSES
    # --------------------------------------------------------

    carpetas_paises = [
        carpeta
        for carpeta in raiz.iterdir()
        if carpeta.is_dir()
    ]

    for carpeta_pais in carpetas_paises:

        nombre_pais = (
            carpeta_pais.name.strip()
        )

        pais = buscar_pais(
            nombre_pais,
            paises,
        )

        if pais is None:

            print(
                "\n⏭ País no encontrado "
                "en la BD:"
            )

            print(
                f"   {nombre_pais}"
            )

            total_omitidas += 1

            continue

        pais_id = pais.get("id")

        # ----------------------------------------------------
        # OPERADORES
        # ----------------------------------------------------

        carpetas_operadores = [
            carpeta
            for carpeta
            in carpeta_pais.iterdir()
            if carpeta.is_dir()
        ]

        for carpeta_to in (
            carpetas_operadores
        ):

            nombre_operador = (
                carpeta_to.name.strip()
            )

            if (
                solo_operador
                and normalizar_nombre(nombre_operador)
                != normalizar_nombre(solo_operador)
            ):
                continue

            operador = buscar_operador(
                nombre_operador,
                pais_id,
                operadores_backend,
            )

            operador_id = None

            if operador is None:

                if dry_run:

                    print(
                        "\n🆕 DRY-RUN "
                        "operador nuevo:"
                    )

                    print(
                        f"   País: "
                        f"{pais.get('nombre')}"
                    )

                    print(
                        f"   Operador: "
                        f"{nombre_operador}"
                    )

                    print(
                        "   Se crearía "
                        "automáticamente."
                    )

                else:

                    operador = crear_operador(
                        api_base,
                        pais_id,
                        nombre_operador,
                    )

                    if operador is None:

                        total_error += 1
                        continue

                    operadores_backend.append(
                        operador
                    )

                    operador_id = (
                        operador.get("id")
                    )

                    total_operadores_creados += 1

            else:

                operador_id = (
                    operador.get("id")
                )

            # ------------------------------------------------
            # ACTIVIDADES
            # ------------------------------------------------

            actividades = [
                carpeta
                for carpeta
                in carpeta_to.iterdir()
                if carpeta.is_dir()
            ]

            for carpeta_actividad in (
                actividades
            ):

                nombre_actividad = (
                    carpeta_actividad
                    .name
                    .strip()
                )

                if (
                    solo_actividad
                    and normalizar_nombre(nombre_actividad)
                    != normalizar_nombre(solo_actividad)
                ):
                    continue

                print(
                    "\n"
                    "==========================="
                )

                print(
                    "📁 Revisando actividad"
                )

                print(
                    "==========================="
                )

                print(
                    f"País: "
                    f"{pais.get('nombre')}"
                )

                print(
                    f"Operador: "
                    f"{nombre_operador}"
                )

                print(
                    f"Actividad: "
                    f"{nombre_actividad}"
                )

                archivos = (
                    buscar_documentos_actividad(
                        carpeta_actividad
                    )
                )

                if not archivos:

                    print(
                        "⏭ No se encontraron "
                        "PDF ni DOCX."
                    )

                    total_omitidas += 1
                    continue

                print(
                    f"📄 Documentos encontrados: "
                    f"{len(archivos)}"
                )

                bloques_texto = []

                # --------------------------------------------
                # LEER ARCHIVOS
                # --------------------------------------------

                for archivo in archivos:

                    print(
                        f"   - leyendo "
                        f"{archivo.name}"
                    )

                    texto = (
                        extraer_texto_archivo(
                            archivo
                        )
                    )

                    if not texto.strip():
                        continue

                    if (
                        es_checklist_sin_informacion(
                            texto
                        )
                    ):

                        print(
                            "   - checklist vacío "
                            f"omitido: "
                            f"{archivo.name}"
                        )

                        continue

                    bloques_texto.append(
                        (
                            f"--- "
                            f"{archivo.name} "
                            f"---\n"
                            f"{texto.strip()}"
                        )
                    )

                # --------------------------------------------
                # ANALIZAR
                # --------------------------------------------

                if not bloques_texto:

                    print(
                        "⚠ No se encontró "
                        "contenido útil."
                    )

                    datos_extraidos = {

                        "descripcion":
                            None,

                        "ubicacion":
                            None,

                        "duracion":
                            None,

                        "edadMinima":
                            None,

                        "edadMaxima":
                            None,

                        "idiomas":
                            [],

                        "imagenes":
                            [],

                        "destino":
                            None,

                        "highlights":
                            [],

                        "itinerario":
                            [],

                        "horarios":
                            [],

                        "restricciones":
                            [],
                    }

                else:

                    texto_completo = (
                        "\n\n".join(
                            bloques_texto
                        )
                    )

                    datos_extraidos = (
                        analizar_contenido_actividad(
                            texto_completo,
                            nombre_actividad,
                        )
                    )

                # --------------------------------------------
                # MOSTRAR
                # --------------------------------------------

                print(
                    "\n🔎 Datos detectados:"
                )

                print(
                    "   Descripción:",
                    datos_extraidos[
                        "descripcion"
                    ],
                )

                print(
                    "   Ubicación:",
                    datos_extraidos[
                        "ubicacion"
                    ],
                )

                print(
                    "   Duración:",
                    datos_extraidos[
                        "duracion"
                    ],
                )

                print(
                    "   Edad mínima:",
                    datos_extraidos[
                        "edadMinima"
                    ],
                )

                print(
                    "   Edad máxima:",
                    datos_extraidos[
                        "edadMaxima"
                    ],
                )

                print(
                    "   Idiomas:",
                    datos_extraidos[
                        "idiomas"
                    ],
                )

                print(
                    "   Destino:",
                    datos_extraidos[
                        "destino"
                    ],
                )

                print(
                    "   Highlights:",
                    datos_extraidos[
                        "highlights"
                    ],
                )

                print(
                    "   Itinerario:",
                    datos_extraidos[
                        "itinerario"
                    ],
                )

                print(
                    "   Horarios:",
                    datos_extraidos[
                        "horarios"
                    ],
                )

                print(
                    "   Restricciones:",
                    datos_extraidos[
                        "restricciones"
                    ],
                )

                # --------------------------------------------
                # DRY RUN
                # --------------------------------------------

                if dry_run:

                    print(
                        "\n✅ DRY-RUN"
                    )

                    print(
                        "   No se modificará "
                        "PostgreSQL."
                    )

                    if operador_id is None:

                        print(
                            "   Operador sería "
                            "creado primero."
                        )

                    else:

                        print(
                            f"   Operador ID: "
                            f"{operador_id}"
                        )

                    total_simuladas += 1

                    continue

                # --------------------------------------------
                # PAYLOAD NUEVO
                # --------------------------------------------

                payload_nuevo = {

                    "nombre":
                        nombre_actividad,

                    "descripcion":
                        datos_extraidos[
                            "descripcion"
                        ],

                    "ubicacion":
                        datos_extraidos[
                            "ubicacion"
                        ],

                    "duracion":
                        datos_extraidos[
                            "duracion"
                        ],

                    "edadMinima":
                        datos_extraidos[
                            "edadMinima"
                        ],

                    "edadMaxima":
                        datos_extraidos[
                            "edadMaxima"
                        ],

                    "idiomas":
                        datos_extraidos[
                            "idiomas"
                        ],

                    "imagenes":
                        datos_extraidos[
                            "imagenes"
                        ],

                    "destino":
                        datos_extraidos[
                            "destino"
                        ],

                    "highlights":
                        datos_extraidos[
                            "highlights"
                        ],

                    "itinerario":
                        datos_extraidos[
                            "itinerario"
                        ],

                    "horarios":
                        datos_extraidos[
                            "horarios"
                        ],

                    "restricciones":
                        datos_extraidos[
                            "restricciones"
                        ],

                    "urlOrigen":
                        (
                            "local:"
                            + str(
                                carpeta_actividad
                                .resolve()
                            )
                        ),

                    "operadorTuristico": {
                        "id":
                            operador_id
                    },

                    "activo":
                        True,
                }

                actividad_existente = (
                    obtener_actividad_existente(
                        backend_actividades,
                        operador_id,
                        nombre_actividad,
                    )
                )

                # --------------------------------------------
                # PUT
                # --------------------------------------------

                if actividad_existente:

                    actividad_id = (
                        actividad_existente
                        .get("id")
                    )

                    payload_actualizacion = (
                        construir_payload_actualizacion(
                            payload_nuevo,
                            actividad_existente,
                        )
                    )

                    print(
                        "\n🔄 Actividad "
                        "existente."
                    )

                    print(
                        f"   Actualizando "
                        f"ID: {actividad_id}"
                    )

                    try:

                        respuesta = requests.put(
                            (
                                f"{backend_actividades}/"
                                f"{actividad_id}"
                            ),
                            json=payload_actualizacion,
                            timeout=30,
                        )

                        if respuesta.status_code in (
                            200,
                            201,
                        ):

                            print(
                                "✅ Actualizada "
                                "correctamente."
                            )

                            total_actualizadas += 1

                        else:

                            # Algunos registros antiguos pueden responder 404
                            # cuando el PUT recibe una relación parcial. Si el
                            # registro sí existe, reintentamos una sola vez
                            # partiendo del objeto completo devuelto por GET.
                            if respuesta.status_code == 404:
                                try:
                                    verificacion = requests.get(
                                        f"{backend_actividades}/{actividad_id}",
                                        timeout=20,
                                    )

                                    if verificacion.status_code == 200:
                                        base_completa = verificacion.json()
                                        payload_reintento = dict(base_completa)
                                        payload_reintento.update(payload_actualizacion)

                                        if base_completa.get("operadorTuristico"):
                                            payload_reintento["operadorTuristico"] = (
                                                base_completa["operadorTuristico"]
                                            )

                                        reintento = requests.put(
                                            f"{backend_actividades}/{actividad_id}",
                                            json=payload_reintento,
                                            timeout=30,
                                        )

                                        if reintento.status_code in (200, 201):
                                            print(
                                                "✅ Actualizada correctamente "
                                                "(reintento con entidad completa)."
                                            )
                                            total_actualizadas += 1
                                            continue

                                        respuesta = reintento

                                except requests.RequestException as error_reintento:
                                    print(
                                        f"   ⚠ Reintento no disponible: "
                                        f"{error_reintento}"
                                    )

                            print(
                                "❌ Error "
                                "actualizando:"
                            )

                            print(
                                f"   HTTP "
                                f"{respuesta.status_code}"
                            )

                            cuerpo = respuesta.text[:1000]
                            if cuerpo:
                                print(cuerpo)

                            print(
                                "   Campos enviados:",
                                sorted(payload_actualizacion.keys()),
                            )

                            print(
                                "   Operador enviado:",
                                payload_actualizacion.get("operadorTuristico"),
                            )

                            total_error += 1

                    except requests.RequestException as error:

                        print(
                            "❌ Error PUT:"
                        )

                        print(
                            f"   {error}"
                        )

                        total_error += 1

                    continue

                # --------------------------------------------
                # POST
                # --------------------------------------------

                try:

                    respuesta = requests.post(
                        backend_actividades,
                        json=payload_nuevo,
                        timeout=30,
                    )

                    if respuesta.status_code in (
                        200,
                        201,
                    ):

                        datos = respuesta.json()

                        print(
                            "\n✅ Actividad "
                            "creada."
                        )

                        print(
                            f"   ID: "
                            f"{datos.get('id')}"
                        )

                        total_creadas += 1

                    else:

                        print(
                            "\n❌ Error "
                            "creando actividad:"
                        )

                        print(
                            f"   HTTP "
                            f"{respuesta.status_code}"
                        )

                        print(
                            respuesta.text[:500]
                        )

                        total_error += 1

                except requests.RequestException as error:

                    print(
                        "\n❌ Error POST:"
                    )

                    print(
                        f"   {error}"
                    )

                    total_error += 1

    # ========================================================
    # RESUMEN
    # ========================================================

    print(
        "\n"
        "==========================="
    )

    print(
        "RESUMEN DE LA EJECUCIÓN"
    )

    print(
        "==========================="
    )

    if dry_run:

        print(
            f"Simuladas: "
            f"{total_simuladas}"
        )

    else:

        print(
            f"Operadores creados: "
            f"{total_operadores_creados}"
        )

        print(
            f"Actividades creadas: "
            f"{total_creadas}"
        )

        print(
            f"Actividades actualizadas: "
            f"{total_actualizadas}"
        )

    print(
        f"Omitidas: "
        f"{total_omitidas}"
    )

    print(
        f"Errores: "
        f"{total_error}"
    )


# ============================================================
# MAIN
# ============================================================

def main():

    parser = argparse.ArgumentParser(
        description=(
            "Extrae actividades desde "
            "PDF/Word, detecta países "
            "y operadores y crea o "
            "actualiza la información."
        )
    )

    parser.add_argument(
        "--raiz",
        required=True,
        help=(
            "Carpeta raíz de "
            "INFORMACIÓN TOUR."
        ),
    )

    parser.add_argument(
        "--api",
        default=(
            "http://localhost:8080"
        ),
        help=(
            "URL base del backend."
        ),
    )

    parser.add_argument(
        "--dry-run",
        action="store_true",
        help=(
            "Prueba todo sin modificar "
            "PostgreSQL."
        ),
    )

    parser.add_argument(
        "--solo-actividad",
        default=None,
        help=(
            "Procesa únicamente una actividad con ese nombre exacto "
            "(comparación normalizada). Útil para pruebas."
        ),
    )

    parser.add_argument(
        "--solo-operador",
        default=None,
        help=(
            "Procesa únicamente un operador con ese nombre exacto "
            "(comparación normalizada). Útil para pruebas."
        ),
    )

    args = parser.parse_args()

    raiz = Path(
        args.raiz
    )

    if (
        not raiz.exists()
        or not raiz.is_dir()
    ):

        print(
            f"❌ La carpeta raíz "
            f"no existe: {raiz}"
        )

        sys.exit(1)

    recorrer_y_procesar(
        raiz=raiz,
        api_base=args.api,
        dry_run=args.dry_run,
        solo_actividad=args.solo_actividad,
        solo_operador=args.solo_operador,
    )


if __name__ == "__main__":
    main()